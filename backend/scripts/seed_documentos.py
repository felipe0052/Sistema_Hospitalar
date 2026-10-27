"""Script para geração dos POPs hospitalares de exemplo em PDF/DOCX e upload via API (S3-T3)."""

from io import BytesIO
import os
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
import httpx

from app.core.config import settings
from app.core.database import Base, SessionLocal, engine
from app.main import app
from app.models.documento import Documento

DIR_RAIZ_SISTEMA = Path(__file__).resolve().parents[2]
DIR_EXEMPLOS_POPS = DIR_RAIZ_SISTEMA / "docs" / "exemplos-pops"

CATALOGO_POPS_EXEMPLO: list[dict[str, Any]] = [
    {
        "arquivo": "pop-001-higienizacao-maos.pdf",
        "formato": "pdf",
        "titulo": "POP-CCIH-001 — Higienização Simples e Antisséptica das Mãos",
        "tipo": "POP",
        "categoria": "Enfermagem",
        "descricao": (
            "Padroniza a técnica dos 5 momentos para a higienização das mãos com água e sabonete "
            "líquido ou preparação alcoólica a 70% em todas as unidades assistenciais do Hospital Decós."
        ),
        "palavras_chave": "higienização, mãos, ccih, infecção hospitalar, biossegurança, enfermagem",
        "secoes": [
            ("1. OBJETIVO", "Prevenir e controlar infecções relacionadas à assistência à saúde (IRAS) por meio da higienização correta das mãos pelos profissionais de saúde."),
            ("2. APLICAÇÃO", "Unidades de Terapia Intensiva (UTI), Pronto-Socorro, Centro Cirúrgico, Pediatria, Internação e Laboratório."),
            ("3. OS 5 MOMENTOS PARA HIGIENIZAÇÃO DAS MÃOS", "1) Antes de tocar o paciente; 2) Antes de realizar procedimento limpo/asséptico; 3) Após risco de exposição a fluidos corporais; 4) Após tocar o paciente; 5) Após tocar superfícies próximas ao paciente."),
            ("4. PROCEDIMENTO TÉCNICO", "Friccionar palmas, dorso das mãos, espaços interdigitais, polegares, articulações e punhos por 20 a 30 segundos com solução alcoólica 70%, ou 40 a 60 segundos com água e sabonete antisséptico."),
        ],
    },
    {
        "arquivo": "pop-002-curativo-acesso-venoso-central.pdf",
        "formato": "pdf",
        "titulo": "POP-ENF-002 — Curativo Simples e Complexo em Acesso Venoso Central",
        "tipo": "POP",
        "categoria": "Enfermagem",
        "descricao": (
            "Estabelece a técnica asséptica e a periodicidade de troca de curativos de cateter "
            "venoso central (CVC) com filme transparente estéril ou gaze na UTI e clínicas."
        ),
        "palavras_chave": "curativo, cateter, acesso venoso central, cvc, uti, enfermagem, assepsia",
        "secoes": [
            ("1. OBJETIVO", "Manter a permeabilidade e estabilidade do cateter venoso central e prevenir infecção primária de corrente sanguínea (IPCS)."),
            ("2. MATERIAIS NECESSÁRIOS", "Luva estéril, máscara cirúrgica, clorexidina alcoólica 0,5%, pacote de curativo estéril, soro fisiológico 0,9% e película transparente semipermeável estéril."),
            ("3. PERIODICIDADE DE TROCA", "Cobertura com gaze estéril: trocar a cada 48 horas. Película transparente estéril: trocar a cada 7 dias ou imediatamente se houver sujidade, umidade ou descolamento."),
            ("4. REGISTRO EM PRONTUÁRIO", "Anotar data, hora, aspecto do óstio de inserção (hiperemia, secreção), tipo de cobertura aplicada e assinatura do enfermeiro responsável."),
        ],
    },
    {
        "arquivo": "pop-003-administracao-medicamentos-sonda.pdf",
        "formato": "pdf",
        "titulo": "POP-ENF-003 — Administração Segura de Medicamentos via Sonda Nasoenteral",
        "tipo": "POP",
        "categoria": "Enfermagem",
        "descricao": (
            "Procedimento operacional para preparo, trituração segura, diluição e administração "
            "de medicamentos por sonda nasogástrica ou nasoenteral com prevenção de obstrução."
        ),
        "palavras_chave": "sonda, medicamentos, nasogástrica, nasoenteral, enfermagem, segurança do paciente",
        "secoes": [
            ("1. OBJETIVO", "Garantir a absorção terapêutica adequada e evitar obstrução da sonda enteral ou incompatibilidades físico-químicas."),
            ("2. DIRETRIZES DE SEGURANÇA", "Conferir os 9 certos da administração de medicamentos. Nunca triturar comprimidos de liberação prolongada (XR/SR) ou revestimento entérico; solicitar avaliação da Farmácia Clínica."),
            ("3. ETAPAS DE EXECUÇÃO", "1) Pausar dieta enteral 15 minutos antes; 2) Testar posicionamento da sonda; 3) Lavar a sonda com 20 mL de água filtrada; 4) Administrar cada medicamento separadamente, lavando com 10 mL entre eles; 5) Lavar ao final com 20 mL de água."),
        ],
    },
    {
        "arquivo": "pop-004-dispensacao-medicamentos-controlados.pdf",
        "formato": "pdf",
        "titulo": "POP-FAR-004 — Dispensação e Rastreabilidade de Medicamentos Controlados (Portaria 344)",
        "tipo": "POP",
        "categoria": "Farmácia",
        "descricao": (
            "Normas de armazenamento em armário chaveado, conferência dupla, escrituração e "
            "rastreabilidade de lote para psicotrópicos e entorpecentes na Farmácia Hospitalar."
        ),
        "palavras_chave": "farmácia, controlados, portaria 344, psicotrópicos, dispensação, lote, rastreabilidade",
        "secoes": [
            ("1. OBJETIVO", "Assegurar o cumprimento rigoroso da Portaria SVS/MS nº 344/1998 na guarda, dispensação e devolução de medicamentos sujeitos a controle especial."),
            ("2. CONFERÊNCIA DUPLA", "Toda separação de medicamentos potencialmente perigosos e controlados exige dupla checagem (técnico de farmácia e farmacêutico clínico) antes da saída para o posto de enfermagem."),
            ("3. CONTROLE DE PSICOBOX E CARROS DE EMERGÊNCIA", "A reposição de ampolas utilizadas em intercorrências ocorre mediante apresentação da prescrição médica e devolução do casco vazio identificado com etiqueta de rastreabilidade."),
        ],
    },
    {
        "arquivo": "pop-005-triagem-classificacao-risco-manchester.pdf",
        "formato": "pdf",
        "titulo": "POP-PS-005 — Triagem e Classificação de Risco no Pronto-Socorro",
        "tipo": "POP",
        "categoria": "Pronto-Socorro",
        "descricao": (
            "Protocolo clínico de acolhimento com classificação de risco e priorização do "
            "atendimento de urgência e emergência segundo sinais de gravidade."
        ),
        "palavras_chave": "triagem, classificação de risco, pronto-socorro, emergência, acolhimento, sinais vitais",
        "secoes": [
            ("1. OBJETIVO", "Identificar rapidamente pacientes em condição crítica ou ameaçadora à vida logo na chegada ao Pronto-Socorro do Hospital Decós."),
            ("2. TEMPOS-ALVO POR PRIORIDADE", "Vermelho (Emergência): atendimento imediato (0 min); Laranja (Muito Urgente): até 10 min; Amarelo (Urgente): até 60 min; Verde (Pouco Urgente): até 120 min; Azul (Não Urgente): até 240 min."),
            ("3. AFERIÇÃO OBRIGATÓRIA", "Aferir pressão arterial, frequência cardíaca, frequência respiratória, temperatura axilar, oximetria de pulso (SpO2), glicemia capilar se indicado e escala de dor."),
        ],
    },
    {
        "arquivo": "pop-006-coleta-transporte-amostras-biologicas.docx",
        "formato": "docx",
        "titulo": "POP-LAB-006 — Coleta, Identificação e Transporte de Amostras Biológicas",
        "tipo": "POP",
        "categoria": "Laboratório",
        "descricao": (
            "Padronização da fase pré-analítica: conferência de identificação dupla à beira-leito, "
            "ordem de tubos de coleta venosa e acondicionamento térmico até a bancada."
        ),
        "palavras_chave": "laboratório, coleta, sangue, hemocultura, amostras, tubos, identificação",
        "secoes": [
            ("1. OBJETIVO", "Prevenir erros pré-analíticos como hemólise, coagulação da amostra ou troca de identificação de pacientes."),
            ("2. IDENTIFICAÇÃO SEGURA À BEIRA-LEITO", "Conferir nome completo e data de nascimento na pulseira do paciente antes da punção e etiquetar os tubos na presença do paciente."),
            ("3. ORDEM DOS TUBOS DE COLETA A VÁCUO", "1º Frascos de hemocultura; 2º Tampa azul (citrato); 3º Tampa vermelha/amarela (soro/gel); 4º Tampa verde (heparina); 5º Tampa roxa (EDTA); 6º Tampa cinza (fluoreto)."),
        ],
    },
    {
        "arquivo": "doc-007-fluxo-acolhimento-internacao-recepcao.docx",
        "formato": "docx",
        "titulo": "Protocolo Institucional de Acolhimento, Internação e Altas na Recepção Central",
        "tipo": "DOCUMENTO",
        "categoria": "Administrativo",
        "descricao": (
            "Diretrizes operacionais de atendimento humanizado na Recepção Central, conferência "
            "documental, emissão de pulseira de identificação e comunicação com os andares."
        ),
        "palavras_chave": "recepção, internação, alta hospitalar, administrativo, atendimento, cadastro",
        "secoes": [
            ("1. OBJETIVO", "Padronizar o atendimento administrativo de admissão eletiva e de urgência, assegurando agilidade e conferência cadastral."),
            ("2. DOCUMENTAÇÃO EXIGIDA NA INTERNAÇÃO", "Documento oficial com foto, CPF, carteirinha do convênio, pedido médico de internação e termo de consentimento institucional assinado."),
            ("3. FLUXO DE COMUNICAÇÃO", "Após emissão da pulseira de identificação branca, acionar o posto de enfermagem do andar de destino pelo ramal correspondente para preparo do leito."),
        ],
    },
]


def _escapar_pdf(texto: str) -> str:
    # Converte para latin-1 compatível com fontes Type1 padrão do PDF (Helvetica)
    codificado = texto.encode("cp1252", errors="replace").decode("latin-1")
    return codificado.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def gerar_bytes_pdf_realista(titulo: str, categoria: str, descricao: str, secoes: list[tuple[str, str]]) -> bytes:
    """Gera um arquivo PDF 1.4 válido e legível em visualizadores PDF (iframe do navegador)."""
    linhas_conteudo: list[str] = [
        "BT",
        "/F1 14 Tf",
        "50 780 Td",
        f"({_escapar_pdf('HOSPITAL DECOS CORPORATIVO - CENTRAL DE NORMAS E POPs')}) Tj",
        "/F1 11 Tf",
        "0 -22 Td",
        f"({_escapar_pdf(titulo[:95])}) Tj",
        "/F2 9 Tf",
        "0 -18 Td",
        f"({_escapar_pdf(f'Categoria: {categoria}   |   Documento Vigente - Qualidade ONA Nivel III')}) Tj",
        "0 -20 Td",
        f"({_escapar_pdf(descricao[:110])}) Tj",
    ]
    if len(descricao) > 110:
        linhas_conteudo.extend(["0 -14 Td", f"({_escapar_pdf(descricao[110:220])}) Tj"])

    for subtitulo, paragrafo in secoes:
        linhas_conteudo.extend(
            [
                "/F1 10 Tf",
                "0 -24 Td",
                f"({_escapar_pdf(subtitulo)}) Tj",
                "/F2 9 Tf",
            ]
        )
        # Quebra parágrafo em linhas de ~95 caracteres
        palavras = paragrafo.split(" ")
        linha_atual = ""
        for p in palavras:
            if len(linha_atual) + len(p) + 1 <= 95:
                linha_atual = f"{linha_atual} {p}".strip()
            else:
                linhas_conteudo.extend(["0 -14 Td", f"({_escapar_pdf(linha_atual)}) Tj"])
                linha_atual = p
        if linha_atual:
            linhas_conteudo.extend(["0 -14 Td", f"({_escapar_pdf(linha_atual)}) Tj"])

    linhas_conteudo.append("ET")
    stream_bytes = "\n".join(linhas_conteudo).encode("latin-1", errors="ignore")

    objetos: list[bytes] = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        (
            b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 4 0 R /F2 5 0 R >> >> /Contents 6 0 R >>\nendobj\n"
        ),
        b"4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>\nendobj\n",
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>\nendobj\n",
        f"6 0 obj\n<< /Length {len(stream_bytes)} >>\nstream\n".encode("ascii")
        + stream_bytes
        + b"\nendstream\nendobj\n",
    ]

    cabecalho = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    corpo = bytearray(cabecalho)
    offsets: list[int] = []
    for obj in objetos:
        offsets.append(len(corpo))
        corpo.extend(obj)

    xref_pos = len(corpo)
    xref_linhas = [f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n"]
    for off in offsets:
        xref_linhas.append(f"{off:010d} 00000 n \n")
    trailer = (
        "".join(xref_linhas)
        + f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    )
    corpo.extend(trailer.encode("ascii"))
    return bytes(corpo)


def gerar_bytes_docx_realista(titulo: str, categoria: str, descricao: str, secoes: list[tuple[str, str]]) -> bytes:
    """Gera um arquivo DOCX válido utilizando python-docx."""
    doc = DocxDocument()
    doc.add_heading("HOSPITAL DECÓS CORPORATIVO — CENTRAL DE DOCUMENTOS & POPs", level=1)
    doc.add_heading(titulo, level=2)
    doc.add_paragraph(f"Categoria: {categoria} | Vigência: Institucional")
    doc.add_paragraph(descricao)
    for subtitulo, texto in secoes:
        doc.add_heading(subtitulo, level=3)
        doc.add_paragraph(texto)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def garantir_arquivos_exemplos_disco(diretorio: Path = DIR_EXEMPLOS_POPS) -> list[tuple[dict[str, Any], Path]]:
    """Gera os 7 arquivos PDF/DOCX de exemplo em docs/exemplos-pops/ e retorna seus caminhos."""
    diretorio.mkdir(parents=True, exist_ok=True)
    gerados: list[tuple[dict[str, Any], Path]] = []
    for item in CATALOGO_POPS_EXEMPLO:
        caminho = diretorio / item["arquivo"]
        if item["formato"] == "pdf":
            conteudo = gerar_bytes_pdf_realista(
                titulo=item["titulo"],
                categoria=item["categoria"],
                descricao=item["descricao"],
                secoes=item["secoes"],
            )
        else:
            conteudo = gerar_bytes_docx_realista(
                titulo=item["titulo"],
                categoria=item["categoria"],
                descricao=item["descricao"],
                secoes=item["secoes"],
            )
        caminho.write_bytes(conteudo)
        gerados.append((item, caminho))
    return gerados


def _enviar_pops_com_cliente(client: Any, base_prefix: str = "/api/v1") -> int:
    arquivos = garantir_arquivos_exemplos_disco()

    email_admin = os.getenv("SEED_ADMIN_EMAIL", "admin@decos.com.br")
    senha_admin = os.getenv("SEED_ADMIN_PASSWORD", "Admin@123")

    res_login = client.post(
        f"{base_prefix}/auth/login",
        json={"email": email_admin, "senha": senha_admin},
    )
    if res_login.status_code != 200:
        raise RuntimeError(
            f"Falha ao autenticar como admin ({email_admin}). Execute primeiro scripts/seed_usuarios.py. "
            f"Status: {res_login.status_code}"
        )

    token = res_login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # Consulta documentos já existentes para evitar duplicidade
    res_existentes = client.get(f"{base_prefix}/documentos?tamanho=100", headers=headers)
    titulos_existentes: set[str] = set()
    if res_existentes.status_code == 200:
        for doc in res_existentes.json().get("itens", []):
            titulos_existentes.add(doc["titulo"].strip().lower())

    inseridos = 0
    for item, caminho in arquivos:
        if item["titulo"].strip().lower() in titulos_existentes:
            continue

        mime = (
            "application/pdf"
            if item["formato"] == "pdf"
            else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        res_up = client.post(
            f"{base_prefix}/documentos",
            headers=headers,
            data={
                "titulo": item["titulo"],
                "tipo": item["tipo"],
                "categoria": item["categoria"],
                "descricao": item["descricao"],
                "palavras_chave": item["palavras_chave"],
            },
            files={"arquivo": (caminho.name, caminho.read_bytes(), mime)},
        )
        if res_up.status_code == 201:
            inseridos += 1
        else:
            raise RuntimeError(f"Erro ao enviar {caminho.name}: {res_up.status_code} {res_up.text}")

    return inseridos


def executar_seed_documentos() -> None:
    """Popula os POPs e documentos de exemplo via API HTTP (ou via TestClient se o servidor não estiver ativo)."""
    if settings.DATABASE_URL.startswith("sqlite"):
        Base.metadata.create_all(bind=engine)

    api_url = os.getenv("API_URL", "http://localhost:8000/api/v1").rstrip("/")
    inseridos = 0
    try:
        with httpx.Client(timeout=5.0) as http_client:
            health = http_client.get(f"{api_url}/health")
            if health.status_code == 200:
                inseridos = _enviar_pops_com_cliente(http_client, base_prefix=api_url)
            else:
                raise httpx.ConnectError("Health check não retornou 200")
    except Exception:
        from fastapi.testclient import TestClient

        with TestClient(app) as local_client:
            inseridos = _enviar_pops_com_cliente(local_client, base_prefix="/api/v1")

    db = SessionLocal()
    try:
        total = db.query(Documento).count()
    finally:
        db.close()

    print(
        f"Seed de documentos concluído: {inseridos} novos POPs/documentos enviados via API "
        f"({total} documentos no banco; arquivos salvos em {DIR_EXEMPLOS_POPS})."
    )


if __name__ == "__main__":
    executar_seed_documentos()
