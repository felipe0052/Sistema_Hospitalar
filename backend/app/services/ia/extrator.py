"""Extração de texto de PDF e DOCX e divisão em trechos (chunking) com sobreposição."""

import re
from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.services.documento_service import MIME_DOCX, MIME_PDF

TAMANHO_TRECHO = 1000
SOBREPOSICAO = 200

_CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def limpar_texto(texto: str) -> str:
    """Remove caracteres de controle e normaliza quebras de linha/espaços."""
    texto = _CONTROLE.sub(" ", texto or "")
    texto = texto.replace("\xa0", " ")
    texto = re.sub(r"-\n(?=\w)", "", texto)  # palavra hifenizada na quebra de linha
    texto = re.sub(r"\s*\n\s*", " ", texto)
    return re.sub(r"[ \t]{2,}", " ", texto).strip()


def extrair_paginas(conteudo: bytes, mime_type: str) -> list[tuple[int, str]]:
    """Retorna [(pagina, texto)] só com páginas que tenham texto. DOCX não tem paginação: tudo na página 1."""
    if mime_type == MIME_PDF:
        leitor = PdfReader(BytesIO(conteudo))
        paginas = [(i, limpar_texto(p.extract_text() or "")) for i, p in enumerate(leitor.pages, start=1)]
    elif mime_type == MIME_DOCX:
        doc = DocxDocument(BytesIO(conteudo))
        partes = [p.text for p in doc.paragraphs]
        for tabela in doc.tables:
            for linha in tabela.rows:
                partes.append(" | ".join(c.text for c in linha.cells))
        paginas = [(1, limpar_texto("\n".join(partes)))]
    else:
        raise ValueError(f"Tipo de arquivo não suportado para extração: {mime_type}")
    return [(pagina, texto) for pagina, texto in paginas if texto]


def dividir_em_trechos(
    paginas: list[tuple[int, str]],
    tamanho: int = TAMANHO_TRECHO,
    sobreposicao: int = SOBREPOSICAO,
) -> list[tuple[int, str]]:
    """Divide cada página em trechos de ~`tamanho` caracteres, com `sobreposicao`, preservando a página."""
    trechos: list[tuple[int, str]] = []
    passo = max(tamanho - sobreposicao, 1)
    for pagina, texto in paginas:
        inicio = 0
        while inicio < len(texto):
            fim = min(inicio + tamanho, len(texto))
            if fim < len(texto):
                corte = texto.rfind(" ", inicio + passo // 2, fim)  # evita cortar palavra no meio
                fim = corte if corte > inicio else fim
            trecho = texto[inicio:fim].strip()
            if trecho:
                trechos.append((pagina, trecho))
            if fim >= len(texto):
                break
            inicio = max(fim - sobreposicao, inicio + 1)
    return trechos
